"""
build_regression77_docx.py

Regenerates "Supplementary Data 1" (the Regression77 figure) using the
corrected, regression-free methodology: for each of the 77 canonical ECA
rules, plot the actual mean evolution trajectory (N_REPEATS independent
random arrangements) for three illustrative initial values (100, 200,
300), and mark the value read directly at generation 60 on each
trajectory -- no line fitting.

Reproduces the original document's layout: a 26-row x 3-column table of
panels in ascending rule-number order (77 rules -> 78 cells, last cell
left empty), so it can be dropped into the manuscript's Supplementary
Data 1 slot in place of the original.

Dependencies: numpy, matplotlib, python-docx  (pip install python-docx)

Usage:
    python build_regression77_docx.py
    python build_regression77_docx.py --n-repeats 200 --out My_Fig.docx
"""
import argparse
import os
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from eca_core import example_trajectories

RULES = [8, 6, 2, 34, 38, 12, 10, 14, 26, 30, 58, 62, 74, 78, 106, 110, 40, 44,
         42, 46, 136, 140, 138, 142, 154, 158, 168, 172, 170, 174, 186, 190,
         130, 134, 162, 166, 234, 238, 202, 206, 24, 28, 56, 60, 152, 156,
         184, 188, 4, 18, 22, 36, 50, 54, 72, 76, 90, 94, 104, 108, 122, 126,
         132, 146, 150, 160, 164, 178, 182, 200, 204, 218, 222, 232, 236,
         250, 254]
RULES_ASCENDING = sorted(RULES)   # matches the original document's panel order

INIT_VALUES = [100, 200, 300]
READ_GEN = 60
COLORS = {100: "#0072B2", 200: "#009E73", 300: "#CC79A7"}  # Okabe-Ito, colour-blind-safe
FONT = "Times New Roman"


def make_panel(rule: int, n_repeats: int, panel_dir: str, rng: np.random.Generator) -> str:
    traj = example_trajectories(rule, INIT_VALUES, n_repeats, rng=rng)  # shape (3, 401)
    fig, ax = plt.subplots(figsize=(4.39, 3.29), dpi=150)
    x = np.arange(traj.shape[1])
    for j, k in enumerate(INIT_VALUES):
        y = traj[j]
        val = y[READ_GEN]
        ax.plot(x, y, color=COLORS[k], lw=1.3, label=f"Initial {k}: {val:.2f}")
        ax.scatter([READ_GEN], [val], color=COLORS[k], zorder=5, s=28,
                   edgecolor="black", linewidth=0.5)
    ax.axvline(READ_GEN, color="gray", lw=0.8, ls="--")
    ax.set_xlim(0, 400)
    ax.set_ylim(0, 400)
    ax.set_xlabel("Generation", fontsize=8)
    ax.set_ylabel("Number of 1s", fontsize=8)
    ax.set_title(f"Rule {rule}", fontsize=11)
    ax.tick_params(labelsize=7)
    ax.legend(fontsize=6, loc="best", framealpha=0.85)
    fig.tight_layout(pad=0.6)
    path = os.path.join(panel_dir, f"rule{rule}.png")
    fig.savefig(path, dpi=150)
    plt.close(fig)
    return path


def build_docx(panel_paths, n_repeats: int, out_path: str):
    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(1)
    section.top_margin = section.bottom_margin = Inches(0.75)

    title = doc.add_paragraph()
    run = title.add_run("Supplementary Data 1 (corrected): Evolution Pattern and Direct-Read Validation")
    run.bold = True
    run.font.size = Pt(15)
    run.font.name = FONT

    p = doc.add_paragraph()
    r = p.add_run(
        "Number of 1s and their evolution patterns over 400 generations for 77 rules: "
        "each panel shows the actual mean trajectory ({n} independent random seedings "
        "per initial value; standard Wolfram elementary-cellular-automaton convention, "
        "zero boundary) for initial values 100, 200, and 300. The dot on each trajectory "
        "marks the value read directly from that trajectory at generation 60 — the "
        "value used in the WeightByInitial lookup table.".format(n=n_repeats)
    )
    r.font.name = FONT
    r.font.size = Pt(11)

    note = doc.add_paragraph()
    note_run = note.add_run(
        "Correction note: the original figure/table was built from a single random "
        "arrangement per initial value (no repetition), fitting a linear regression "
        "across the full 0–400-generation trajectory and reading its value at "
        "generation 60. Independent validation found this global-line fit introduces a "
        "substantial, systematic bias (median 1.9–2.3, up to 48 on the 0–400 scale, "
        "across the 77 rules), because most of the fitted range lies well past "
        "generation 60. This figure has been corrected to read each trajectory directly "
        "at generation 60, with the repeat count increased to remove sampling noise. "
        "See the accompanying Supplementary Note, ‘Validation and Correction of the ECA "
        "Generation-60 Lookup Table (WeightByInitial),’ for the full methodology and "
        "validation results."
    )
    note_run.font.name = FONT
    note_run.font.size = Pt(10.5)
    note_run.italic = True

    n_cols = 3
    n_rows = -(-len(panel_paths) // n_cols)  # ceil division
    table = doc.add_table(rows=n_rows, cols=n_cols)
    for i, path in enumerate(panel_paths):
        r_, c_ = divmod(i, n_cols)
        cell_p = table.cell(r_, c_).paragraphs[0]
        cell_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell_p.add_run().add_picture(path, width=Inches(1.9948), height=Inches(1.496))

    doc.save(out_path)
    print(f"written {out_path}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--n-repeats", type=int, default=100,
                     help="independent random arrangements averaged per panel line (default 100)")
    ap.add_argument("--panel-dir", type=str, default="panels")
    ap.add_argument("--out", type=str, default="Regression77_corrected.docx")
    ap.add_argument("--seed", type=int, default=20260921)
    args = ap.parse_args()

    os.makedirs(args.panel_dir, exist_ok=True)
    rng = np.random.default_rng(args.seed)
    panel_paths = [make_panel(rule, args.n_repeats, args.panel_dir, rng) for rule in RULES_ASCENDING]
    build_docx(panel_paths, args.n_repeats, args.out)


if __name__ == "__main__":
    main()